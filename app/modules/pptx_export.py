"""帧序列 → PPT 幻灯片：每帧一页，图片铺满页面（页面尺寸随首帧宽高比）。"""
from __future__ import annotations

from pathlib import Path

from PySide6.QtCore import QThread, Signal


def _image_size(path: Path) -> tuple[int, int]:
    """读取图片像素尺寸（宽, 高）。优先 python-pptx 内部 Image，失败读 PNG 头。"""
    try:
        from pptx.parts.image import Image
        return Image.from_file(str(path)).size
    except Exception:
        pass
    # PNG 兜底：IHDR 在第 16~24 字节
    with open(path, "rb") as f:
        head = f.read(24)
    if head[:8] == b"\x89PNG\r\n\x1a\n":
        return (int.from_bytes(head[16:20], "big"),
                int.from_bytes(head[20:24], "big"))
    return 1920, 1080


class PptxJob(QThread):
    """后台把帧序列写成 .pptx（每帧一页幻灯片）。"""

    progressed = Signal(int, str)          # 百分比, 状态文本
    done = Signal(bool, str)

    def __init__(self, frames: list[Path], out_file: str, parent=None):
        super().__init__(parent)
        self._frames = frames
        self._out = out_file
        self._stopped = False

    def run(self):  # noqa: D102
        try:
            from pptx import Presentation
            from pptx.util import Emu

            if not self._frames:
                self.done.emit(False, "没有可导出的帧")
                return
            w_px, h_px = _image_size(self._frames[0])
            prs = Presentation()
            # 页面宽固定 13.333 英寸（16:9 默认宽），高按首帧宽高比换算
            prs.slide_width = Emu(12192000)
            prs.slide_height = Emu(int(12192000 * h_px / max(w_px, 1)))
            blank = prs.slide_layouts[6]
            total = len(self._frames)
            for i, f in enumerate(self._frames):
                if self._stopped:
                    self.done.emit(False, "已取消")
                    return
                slide = prs.slides.add_slide(blank)
                slide.shapes.add_picture(str(f), 0, 0,
                                         width=prs.slide_width,
                                         height=prs.slide_height)
                self.progressed.emit((i + 1) * 100 // total,
                                     f"已写入 {i + 1}/{total} 页")
            prs.save(self._out)
            self.done.emit(True, f"PPT 已生成（{total} 页）")
        except ImportError:
            self.done.emit(False, "缺少 python-pptx，请执行 py -m pip install python-pptx")
        except Exception as exc:  # noqa: BLE001
            self.done.emit(False, str(exc))

    def stop(self):  # noqa: D102
        self._stopped = True
